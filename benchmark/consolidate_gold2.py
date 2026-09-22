#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Consolida las hojas anotadas del gold limpio (gold2) en un TSV evaluable.

Este es el conjunto que corrige los tres sesgos del gold original: cero solape con el
entrenamiento, distribucion de frecuencias igual a la del corpus y ningun checkpoint
seleccionado sobre el.

La consolidacion es deliberadamente estricta. Una etiqueta mal interpretada aqui contamina
todas las metricas del articulo, asi que ante cualquier ambiguedad se EXCLUYE la fila y se
reporta, en vez de adivinar. Casos contemplados:

  revision = OK          -> se acepta la placa sugerida
  revision = CORREGIR    -> se toma placa_correcta (obligatoria)
  revision = ILEGIBLE    -> se descarta (no hay verdad-terreno posible)
  revision = <una placa> -> el anotador escribio la correccion en la columna equivocada;
                            se interpreta segun coincida o no con la sugerida, y se registra
  revision = OK + correccion distinta -> CONTRADICTORIO: se descarta y se reporta
  revision vacia         -> sin anotar: se descarta y se reporta

Uso:
    python consolidate_gold2.py
"""
import argparse, csv, glob, os, re
from collections import Counter
from pathlib import Path

import openpyxl

from bench_utils import norm_plate

VALIDAS = {"OK", "CORREGIR", "ILEGIBLE"}
COL = {"id": 1, "sugerida": 3, "revision": 4, "correcta": 5, "recorte": 6, "motivo": 7, "tipo": 8}


def leer_hoja(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    filas = []
    for r in range(5, ws.max_row + 1):
        gid = ws.cell(r, COL["id"]).value
        if not gid:
            continue
        filas.append({
            "gid": str(gid).strip(),
            "sugerida": norm_plate(ws.cell(r, COL["sugerida"]).value or ""),
            "revision": (str(ws.cell(r, COL["revision"]).value).strip()
                         if ws.cell(r, COL["revision"]).value is not None else ""),
            "correcta": norm_plate(ws.cell(r, COL["correcta"]).value or ""),
            "recorte": (ws.cell(r, COL["recorte"]).value or "").strip() or "",
            "tipo": (ws.cell(r, COL["tipo"]).value or "").strip() or "",
            "hoja": os.path.basename(path),
            "fila": r,
        })
    return filas


def resolver(f):
    """Devuelve (placa, estado, nota). placa='' cuando la fila no es utilizable."""
    rev, sug, cor = f["revision"], f["sugerida"], f["correcta"]

    if not rev:
        return "", "SIN_ANOTAR", "revision vacia"

    if rev == "ILEGIBLE":
        return "", "ILEGIBLE", ""

    if rev == "OK":
        if cor and cor != sug:
            return "", "CONTRADICTORIO", f"marcada OK pero con correccion {cor}"
        return sug, "OK", ""

    if rev == "CORREGIR":
        if not cor:
            return "", "INCOMPLETA_CORREGIR", "CORREGIR sin placa_correcta"
        return cor, "CORREGIDA", ""

    # La revision no es una de las tres etiquetas: puede ser una placa escrita
    # en la columna equivocada.
    como_placa = norm_plate(rev)
    if como_placa:
        if como_placa == sug:
            return sug, "OK", f"placa en columna revision, coincide con la sugerida ({rev})"
        return como_placa, "CORREGIDA", f"correccion escrita en columna revision ({rev})"

    return "", "REVISION_INVALIDA", f"valor no reconocido: {rev!r}"


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dirs", default="../docs/verificacion_gold2,../docs/verificacion_gold2_noche")
    ap.add_argument("--sel", default="gold2_selection.csv,gold2_noche_selection.csv")
    ap.add_argument("--out", default="data/rec_labels_gold2.tsv")
    ap.add_argument("--meta", default="data/rec_labels_gold2_meta.csv")
    ap.add_argument("--incluir-incompletas", action="store_true",
                    help="incluye los recortes marcados INCOMPLETA (por defecto se excluyen)")
    args = ap.parse_args()

    # --- Atributos de cada recorte, desde las selecciones originales
    atributos = {}
    for sel in args.sel.split(","):
        p = Path(sel.strip())
        if not p.exists():
            continue
        with open(p, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                atributos[row["fn"]] = row

    total = Counter()
    registros, incidencias = [], []

    for d in args.dirs.split(","):
        d = Path(d.strip())
        mapping_path = d / "mapping.csv"
        if not mapping_path.exists():
            print(f"AVISO: falta {mapping_path}, se omite {d}")
            continue
        with open(mapping_path, encoding="utf-8-sig") as f:
            mapping = {r["id"]: r for r in csv.DictReader(f)}

        for hoja in sorted(glob.glob(str(d / "lote_*.xlsx"))):
            for f in leer_hoja(hoja):
                m = mapping.get(f["gid"])
                if m is None:
                    incidencias.append((f, "SIN_MAPPING", "id ausente en mapping.csv"))
                    total["SIN_MAPPING"] += 1
                    continue
                fn = m["nombre_archivo"]
                placa, estado, nota = resolver(f)
                total[estado] += 1
                if nota:
                    incidencias.append((f, estado, nota))
                if not placa:
                    continue
                if f["recorte"] == "INCOMPLETA" and not args.incluir_incompletas:
                    total["EXCLUIDA_INCOMPLETA"] += 1
                    continue
                attr = atributos.get(fn, {})
                registros.append({
                    "path": attr.get("path", ""),
                    "fn": fn,
                    "placa": placa,
                    "sugerida": f["sugerida"],
                    "corregida": int(placa != f["sugerida"]),
                    "recorte": f["recorte"],
                    "tipo_placa": f["tipo"],
                    "luz": attr.get("day_night", ""),
                    "area_px2": attr.get("area_px2", ""),
                    "frames_de_la_placa": attr.get("frames_de_la_placa", ""),
                    "split_origen": attr.get("split_origen", ""),
                    "lote": f["hoja"],
                    "origen_dir": d.name,
                })

    faltan_path = [r for r in registros if not r["path"] or not Path(r["path"]).exists()]
    if faltan_path:
        print(f"AVISO: {len(faltan_path)} registros sin imagen localizable; se excluyen")
        registros = [r for r in registros if r["path"] and Path(r["path"]).exists()]

    # --- Salidas
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n")
        for r in registros:
            w.writerow([r["path"], r["placa"]])

    campos = list(registros[0].keys()) if registros else []
    with open(args.meta, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campos)
        w.writeheader()
        w.writerows(registros)

    # --- Reporte
    print(f"Filas leidas por estado:")
    for k, v in total.most_common():
        print(f"  {k:22s} {v:5d}")
    print(f"\nGOLD2 utilizable: {len(registros)} placas -> {out}")

    corregidas = sum(r["corregida"] for r in registros)
    print(f"  corregidas por el anotador: {corregidas} ({100*corregidas/len(registros):.1f} %)")
    for campo in ("luz", "origen_dir", "tipo_placa"):
        c = Counter(r[campo] for r in registros)
        print(f"  {campo}: " + "  ".join(f"{k or '(vacio)'}={v}" for k, v in c.most_common(6)))

    if incidencias:
        print(f"\nIncidencias que requieren atencion ({len(incidencias)}):")
        for f, estado, nota in incidencias:
            print(f"  {f['origen_dir'] if 'origen_dir' in f else '':0s}{f['hoja']:14s} fila {f['fila']:4d} "
                  f"{f['gid']:6s} [{estado}] {nota}")


if __name__ == "__main__":
    main()
